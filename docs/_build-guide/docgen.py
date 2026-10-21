# -*- coding: utf-8 -*-
"""The Word building blocks every technical document shares: the palette, Arabic text with its Latin runs
left-to-right, headings, tables, notes, the cover, the running header and footer, and the right-to-left pass
before saving. A builder calls new_document(), writes with these, and ends with finish(path).

Used by build_schema.py (هيكلية قاعدة البيانات) and build_rules.py (قواعد العمل والمصطلحات)."""
import re, datetime
from pathlib import Path

ROOT = Path(__file__).parent
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "itf Qomra Arabic"; CODE = "Consolas"
GREEN = RGBColor(0x00, 0x59, 0x4F); GOLD = RGBColor(0xAD, 0x9E, 0x6E); MAROON = RGBColor(0x67, 0x21, 0x46)
INK = RGBColor(0x02, 0x15, 0x26); SOFT = RGBColor(0x55, 0x5F, 0x6B); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LOGO = ROOT.parent.parent / "public" / "icons" / "icon-512.png"
_LTR = re.compile(r"[A-Za-z0-9_\-./:<>{}()\[\]@'=+*#|,]+(?:[ ][A-Za-z0-9_\-./:<>{}()\[\]@'=+*#|,]+)*")
doc = None


def rpr(run):
    r = run._r.get_or_add_rPr()
    return r


def set_font(run, size, bold=False, color=INK, font=FONT, rtl=True):
    r = rpr(run)
    rf = r.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); r.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), font)
    run.font.size = Pt(size)
    sz = OxmlElement("w:szCs"); sz.set(qn("w:val"), str(int(size * 2))); r.append(sz)
    run.font.bold = bold
    if bold:
        b = OxmlElement("w:bCs"); r.append(b)
    run.font.color.rgb = color
    if rtl:
        r.append(OxmlElement("w:rtl"))
        lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "ar-SY"); r.append(lang)


def add_text(p, text, size=11, bold=False, color=INK):
    """Arabic text with its Latin/code stretches run left-to-right"""
    pos = 0
    for m in _LTR.finditer(text):
        tok = m.group(0)
        if not re.search(r"[A-Za-z]", tok): continue  # digits alone sit right in Arabic text
        start = m.start() + (len(tok) - len(tok.lstrip(":,.(")))
        end = m.end() - (len(tok) - len(tok.rstrip(":,.)")))
        if end <= start: continue
        if start > pos:
            r = p.add_run(text[pos:start]); set_font(r, size, bold, color)
        r = p.add_run(text[start:end]); set_font(r, size, bold, color, rtl=False)
        pos = end
    if pos < len(text):
        r = p.add_run(text[pos:]); set_font(r, size, bold, color)


def add_code(p, text, size=9.5, bold=False, color=INK):
    r = p.add_run(text); set_font(r, size, bold, color, font=CODE, rtl=False)


def bidi(p, align=WD_ALIGN_PARAGRAPH.LEFT):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.insert(0, OxmlElement("w:bidi"))
    p.alignment = align
    return p


def para(text="", size=11.5, bold=False, color=INK, after=6, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.25, container=None):
    p = (container or doc).add_paragraph()
    bidi(p, align)
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
    if text: add_text(p, text, size, bold, color)
    return p


def heading(text, level, page_break=False, before=12):
    p = doc.add_paragraph(style=f"Heading {level}")
    bidi(p)
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(6); pf.keep_with_next = True
    if page_break: pf.page_break_before = True
    add_text(p, text, {1: 22, 2: 14.5, 3: 12.5}[level], True, {1: GREEN, 2: MAROON, 3: GREEN}[level])
    return p


def bullets(items, size=11, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style); bidi(p)
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.2
        add_text(p, it, size)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)


def margins(cell, v=50, h=90):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for k, val in (("top", v), ("bottom", v), ("start", h), ("end", h)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def table(headers, rows, widths, head_fill="00594F", size=9.5, code_cols=(), zebra=True, header_size=10):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    tblPr.append(OxmlElement("w:bidiVisual"))
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "D9C89E"); b.append(el)
    tblPr.append(b)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    t.autofit = False
    W = [Cm(w) for w in widths]
    for i, (c, h) in enumerate(zip(t.rows[0].cells, headers)):
        c.width = W[i]; shade(c, head_fill); margins(c)
        p = c.paragraphs[0]; bidi(p); add_text(p, h, header_size, True, WHITE)
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); trPr.append(th)
    for ri, row in enumerate(rows):
        for ci, (c, v) in enumerate(zip(t.rows[ri + 1].cells, row)):
            c.width = W[ci]; margins(c)
            if zebra and ri % 2: shade(c, "F7F4EF")
            p = c.paragraphs[0]; bidi(p); p.paragraph_format.space_after = Pt(0)
            if ci in code_cols: add_code(p, str(v), size - 0.5)
            else: add_text(p, str(v), size)
    para("", after=4)
    return t


def callout(text, fill="EAF4F2", border="8FC7BE", label="ملاحظة", color=GREEN):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr; tblPr.append(OxmlElement("w:bidiVisual"))
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), border); b.append(el)
    tblPr.append(b)
    c = t.rows[0].cells[0]; shade(c, fill); margins(c, 80, 140)
    p = c.paragraphs[0]; bidi(p)
    add_text(p, label + ": ", 10.5, True, color); add_text(p, text, 10.5)
    para("", after=3)


def field(p, instr, placeholder=""):
    for kind, extra in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        r = p.add_run()
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = extra; r._r.append(it)
        elif kind == "text":
            r.text = extra; set_font(r, 10, False, SOFT)
        else:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r._r.append(f)


def new_document():
    """A blank A4 document with the platform's fonts, ready for the cover"""
    global doc
    doc = Document()
    s0 = doc.sections[0]; s0.page_width, s0.page_height = Cm(21), Cm(29.7)
    for s in doc.sections: s.left_margin = s.right_margin = Cm(1.8); s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(11)
    for lvl, sz, col in ((1, 22, GREEN), (2, 14.5, MAROON), (3, 12.5, GREEN)):
        hs = doc.styles[f"Heading {lvl}"]; hs.font.name = FONT; hs.font.size = Pt(sz); hs.font.color.rgb = col; hs.font.bold = True
    return doc


def cover(title, subtitle, tagline, stats, note):
    """The cover page: logo, the platform's name, the title, and a row of figures [(label, value), ...]"""
    para("", after=36)
    p = doc.add_paragraph(); bidi(p, WD_ALIGN_PARAGRAPH.CENTER); p.add_run().add_picture(str(LOGO), width=Cm(4))
    para("إدارة الحج والعمرة السورية", 14, True, GOLD, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("المنصة الوطنية للحج", 18, True, GREEN, after=30, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(title, 40, True, GREEN, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(subtitle, 20, True, MAROON, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(tagline, 13, False, SOFT, after=36, align=WD_ALIGN_PARAGRAPH.CENTER)
    t = doc.add_table(rows=1, cols=len(stats)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    for c, (a, b) in zip(t.rows[0].cells, stats):
        shade(c, "F7F4EF"); margins(c, 120, 120)
        q = c.paragraphs[0]; bidi(q, WD_ALIGN_PARAGRAPH.CENTER); add_text(q, a, 10, True, GOLD)
        q2 = c.add_paragraph(); bidi(q2, WD_ALIGN_PARAGRAPH.CENTER); add_text(q2, b, 12, True, GREEN)
    para("", after=50)
    para(note, 10, False, SOFT, align=WD_ALIGN_PARAGRAPH.CENTER)


def body_section(title):
    """A new section after the cover, with the running header, the page number, and pages counted from 1"""
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.left_margin = body.right_margin = Cm(1.8); body.top_margin = Cm(1.8); body.bottom_margin = Cm(1.6)
    body.header.is_linked_to_previous = False; body.footer.is_linked_to_previous = False
    hp = body.header.paragraphs[0]; bidi(hp, WD_ALIGN_PARAGRAPH.CENTER)
    add_text(hp, title + " ", 9.5, True, GREEN); add_text(hp, "— المنصة الوطنية للحج — موسم 1448هـ", 9.5, False, SOFT)
    fp = body.footer.paragraphs[0]; bidi(fp, WD_ALIGN_PARAGRAPH.CENTER)
    add_text(fp, "صفحة ", 9.5, False, SOFT); field(fp, "PAGE", "1"); add_text(fp, "   |   إدارة الحج والعمرة السورية", 9, False, SOFT)
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); body._sectPr.append(pg)
    return body


def toc():
    heading("فهرس المحتويات", 1, page_break=True, before=0)
    tp = doc.add_paragraph(); bidi(tp)
    field(tp, 'TOC \\o "1-2" \\h \\z \\u', "اضغط بزر الفأرة الأيمن هنا ثم «تحديث الحقل» لعرض الفهرس.")


def finish(out_path):
    """Right-to-left everywhere, table properties in the order Word expects, then save"""
    dd = doc.styles.element.find(qn("w:docDefaults"))
    ppd = dd.find(qn("w:pPrDefault"))
    if ppd is None: ppd = OxmlElement("w:pPrDefault"); dd.append(ppd)
    ppr = ppd.find(qn("w:pPr"))
    if ppr is None: ppr = OxmlElement("w:pPr"); ppd.append(ppr)
    if ppr.find(qn("w:bidi")) is None: ppr.insert(0, OxmlElement("w:bidi"))
    from docx.enum.style import WD_STYLE_TYPE
    for nm in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "toc 1", "toc 2"]:
        try: s = doc.styles[nm]
        except KeyError: s = doc.styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
        pPr = s.element.get_or_add_pPr()
        if pPr.find(qn("w:bidi")) is None: pPr.insert(0, OxmlElement("w:bidi"))
        r = s.element.get_or_add_rPr()
        rf = r.find(qn("w:rFonts"))
        if rf is None: rf = OxmlElement("w:rFonts"); r.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
        if nm.startswith("toc"):
            s.base_style = doc.styles["Normal"]; s.font.size = Pt(11 if nm == "toc 1" else 10); s.font.bold = nm == "toc 1"
    for sct in doc.sections:
        sp = sct._sectPr
        if sp.find(qn("w:bidi")) is None: sp.append(OxmlElement("w:bidi"))
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); doc.settings.element.append(uf)
    _TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
    for tb in doc.tables:
        tp_ = tb._tbl.tblPr; kids = list(tp_)
        for e in kids: tp_.remove(e)
        for e in sorted(kids, key=lambda e: _TBLPR.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in _TBLPR else 99): tp_.append(e)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
