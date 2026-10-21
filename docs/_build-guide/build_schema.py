# -*- coding: utf-8 -*-
"""The database design document and its SQL, from one source (docs/_build-guide/schema):

    python build_schema.py <out dir>          → <out>/هيكلية-قاعدة-البيانات.docx and <out>/schema.sql
    python build_schema.py --check            → only validates the model

The model is checked first (foreign keys and their types, primary keys, the columns of every key), and
nothing is written while it has an error. Then finalize.ps1 updates the table of contents and prints the PDF."""
import sys, re, copy, hashlib, datetime
from pathlib import Path
from collections import defaultdict

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT))
from schema import MODULES, META, INTRO, HOW_TO_READ, HOW_ITEMS, CONVENTIONS, DEMO_GAPS

TABLES = [t for m in MODULES for t in m.tables]
BY_NAME = {t.name: t for t in TABLES}

# MySQL and MariaDB refuse a CHECK on a column whose foreign key has a referential action (SET NULL / CASCADE):
# such a column deletes RESTRICT. Updates RESTRICT everywhere: surrogate ids and reference codes never change.
for _t in TABLES:
    _exprs = " ".join([c.ck for c in _t.cols if c.ck] + list(_t.checks))
    for _c in _t.cols:
        if _c.fk and re.search(rf"\b{re.escape(_c.name)}\b", _exprs): _c.od = "RESTRICT"


# ═══════════════ validation ═══════════════
def base_type(t):
    t = t.upper()
    if t.startswith("VARCHAR"): return "VARCHAR"
    if t.startswith("CHAR"): return "CHAR"
    return t


def effective_pk(t):
    """The declared primary key, else the first unique key whose columns are all NOT NULL"""
    pk = t.pk_cols()
    if pk: return pk
    cols = {c.name: c for c in t.cols}
    for u in t.uq:
        if all(n in cols and not cols[n].null for n in u): return list(u)
    for c in t.cols:
        if c.uq and not c.null: return [c.name]
    return None


def validate():
    errors, warns = [], []
    seen = set()
    for t in TABLES:
        if t.name in seen: errors.append(f"duplicate table {t.name}")
        seen.add(t.name)
        if len(t.name) > 64: errors.append(f"{t.name}: name too long")
        names = [c.name for c in t.cols]
        if len(names) != len(set(names)): errors.append(f"{t.name}: duplicate column {[n for n in names if names.count(n) > 1]}")
        cols = {c.name: c for c in t.cols}
        for key in list(t.uq) + list(t.idx):
            for n in key:
                if n not in cols: errors.append(f"{t.name}: key column {n} missing")
        if not effective_pk(t): errors.append(f"{t.name}: no primary key (add pk() or a NOT NULL unique key)")
        for c in t.cols:
            if c.pk and c.null: errors.append(f"{t.name}.{c.name}: primary key column is nullable")
            if c.ai and "INT" not in c.type.upper(): errors.append(f"{t.name}.{c.name}: AUTO_INCREMENT on a non-integer")
            tgt = c.fk_target()
            if not tgt: continue
            tt, tc = tgt
            if tt not in BY_NAME: errors.append(f"{t.name}.{c.name}: FK to missing table {tt}"); continue
            tcol = {x.name: x for x in BY_NAME[tt].cols}.get(tc)
            if not tcol: errors.append(f"{t.name}.{c.name}: FK to missing column {tt}.{tc}"); continue
            if base_type(tcol.type) != base_type(c.type): errors.append(f"{t.name}.{c.name} {c.type} ≠ {tt}.{tc} {tcol.type}")
            elif tcol.type.upper() != c.type.upper(): warns.append(f"{t.name}.{c.name} {c.type} vs {tt}.{tc} {tcol.type}")
            if not (tcol.pk or tcol.uq or [tc] == effective_pk(BY_NAME[tt]) or any(list(u)[:1] == [tc] for u in BY_NAME[tt].uq)):
                errors.append(f"{t.name}.{c.name}: FK target {tt}.{tc} is not a key")
            if c.od == "SET NULL" and not c.null: errors.append(f"{t.name}.{c.name}: ON DELETE SET NULL on a NOT NULL column")
    return errors, warns


# ═══════════════ SQL ═══════════════
def ident_name(prefix, table, cols):
    name = f"{prefix}_{table}_{'_'.join(cols)}"
    if len(name) <= 64: return name
    h = hashlib.md5(name.encode()).hexdigest()[:8]
    return f"{prefix}_{table[:40]}_{h}"


def sql_str(s, limit):
    s = s.replace("\\", "\\\\").replace("'", "''").replace("\n", " ")
    if len(s) > limit: s = s[:limit - 1] + "…"
    return f"'{s}'"


def col_sql(c, pk_auto):
    parts = [f"`{c.name}`", c.type]
    parts.append("NULL" if c.null else "NOT NULL")
    if c.ai and pk_auto: parts.append("AUTO_INCREMENT")
    if c.d is not None: parts.append(f"DEFAULT {c.d}")
    if c.onupdate: parts.append("ON UPDATE CURRENT_TIMESTAMP(3)")
    parts.append(f"COMMENT {sql_str(c.desc, 1000)}")
    return " ".join(parts)


def build_sql():
    out = []
    w = out.append
    w("-- ═══════════════════════════════════════════════════════════════════════════")
    w(f"-- {META['title']} — {META['subtitle']}")
    w("-- المنصة الوطنية للحج — إدارة الحج والعمرة السورية")
    w(f"-- MySQL 8.0.16+ / MariaDB 10.6+   |   {len(TABLES)} جدولاً   |   مولَّد من docs/_build-guide/schema في {datetime.date.today():%Y-%m-%d}")
    w("-- لا تعدّل هذا الملف يدوياً: عدّل النموذج ثم شغّل build_schema.py")
    w("-- ═══════════════════════════════════════════════════════════════════════════")
    w("")
    w("SET NAMES utf8mb4;")
    w("SET FOREIGN_KEY_CHECKS = 0;")
    w("")
    fks = []
    for m in MODULES:
        w(f"-- ───────────────────────────────────────────────────────────────")
        w(f"-- {m.title}")
        w(f"-- ───────────────────────────────────────────────────────────────")
        w("")
        for t in m.tables:
            lines = []
            pk = effective_pk(t)
            for c in t.cols:
                lines.append("  " + col_sql(c, c.pk))
            lines.append(f"  PRIMARY KEY ({', '.join(f'`{n}`' for n in pk)})")
            uqs = [(c.name,) for c in t.cols if c.uq and [c.name] != pk] + [u for u in t.uq if list(u) != pk]
            for u in uqs:
                lines.append(f"  UNIQUE KEY `{ident_name('uq', t.name, u)}` ({', '.join(f'`{n}`' for n in u)})")
            for ix in t.idx:
                lines.append(f"  KEY `{ident_name('ix', t.name, ix)}` ({', '.join(f'`{n}`' for n in ix)})")
            n_ck = 0
            for c in t.cols:
                if c.ck:
                    n_ck += 1
                    lines.append(f"  CONSTRAINT `{ident_name('ck', t.name, [str(n_ck)])}` CHECK ({c.ck})")
            for ck in t.checks:
                n_ck += 1
                lines.append(f"  CONSTRAINT `{ident_name('ck', t.name, [str(n_ck)])}` CHECK ({ck})")
            w(f"CREATE TABLE `{t.name}` (")
            w(",\n".join(lines))
            w(f") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT={sql_str(t.title + ' — ' + t.desc, 1800)};")
            w("")
            for c in t.cols:
                tgt = c.fk_target()
                if tgt: fks.append((t.name, c, tgt))
    w("-- ───────────────────────────────────────────────────────────────")
    w(f"-- العلاقات ({len(fks)} مفتاحاً أجنبياً)")
    w("-- ───────────────────────────────────────────────────────────────")
    w("")
    for tname, c, (tt, tc) in fks:
        w(f"ALTER TABLE `{tname}` ADD CONSTRAINT `{ident_name('fk', tname, [c.name])}` FOREIGN KEY (`{c.name}`) REFERENCES `{tt}` (`{tc}`) ON DELETE {c.od} ON UPDATE RESTRICT;")
    w("")
    w("SET FOREIGN_KEY_CHECKS = 1;")
    return "\n".join(out) + "\n", len(fks)


# ═══════════════ DOCX ═══════════════
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "itf Qomra Arabic"; CODE = "Consolas"
GREEN = RGBColor(0x00, 0x59, 0x4F); GOLD = RGBColor(0xAD, 0x9E, 0x6E); MAROON = RGBColor(0x67, 0x21, 0x46)
INK = RGBColor(0x02, 0x15, 0x26); SOFT = RGBColor(0x55, 0x5F, 0x6B); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LOGO = ROOT.parent.parent / "public" / "icons" / "icon-512.png"
_LTR = re.compile(r"[A-Za-z0-9_\-./:<>{}()\[\]@'=+*#|,]+(?:[ ][A-Za-z0-9_\-./:<>{}()\[\]@'=+*#|,]+)*")
doc = None


def rpr(run):
    r = run._r.get_or_add_rPr()
    return r


def set_font(run, size, bold=False, color=INK, font=FONT, rtl=True):
    r = rpr(run)
    rf = r.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); r.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), font)
    run.font.size = Pt(size)
    sz = OxmlElement("w:szCs"); sz.set(qn("w:val"), str(int(size * 2))); r.append(sz)
    run.font.bold = bold
    if bold:
        b = OxmlElement("w:bCs"); r.append(b)
    run.font.color.rgb = color
    if rtl:
        r.append(OxmlElement("w:rtl"))
        lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "ar-SY"); r.append(lang)


def add_text(p, text, size=11, bold=False, color=INK):
    """Arabic text with its Latin/code stretches run left-to-right"""
    pos = 0
    for m in _LTR.finditer(text):
        tok = m.group(0)
        if not re.search(r"[A-Za-z]", tok): continue  # digits alone sit right in Arabic text
        start = m.start() + (len(tok) - len(tok.lstrip(":,.(")))
        end = m.end() - (len(tok) - len(tok.rstrip(":,.)")))
        if end <= start: continue
        if start > pos:
            r = p.add_run(text[pos:start]); set_font(r, size, bold, color)
        r = p.add_run(text[start:end]); set_font(r, size, bold, color, rtl=False)
        pos = end
    if pos < len(text):
        r = p.add_run(text[pos:]); set_font(r, size, bold, color)


def add_code(p, text, size=9.5, bold=False, color=INK):
    r = p.add_run(text); set_font(r, size, bold, color, font=CODE, rtl=False)


def bidi(p, align=WD_ALIGN_PARAGRAPH.LEFT):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.insert(0, OxmlElement("w:bidi"))
    p.alignment = align
    return p


def para(text="", size=11.5, bold=False, color=INK, after=6, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.25, container=None):
    p = (container or doc).add_paragraph()
    bidi(p, align)
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line
    if text: add_text(p, text, size, bold, color)
    return p


def heading(text, level, page_break=False, before=12):
    p = doc.add_paragraph(style=f"Heading {level}")
    bidi(p)
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(6); pf.keep_with_next = True
    if page_break: pf.page_break_before = True
    add_text(p, text, {1: 22, 2: 14.5, 3: 12.5}[level], True, {1: GREEN, 2: MAROON, 3: GREEN}[level])
    return p


def bullets(items, size=11, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style); bidi(p)
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.2
        add_text(p, it, size)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)


def margins(cell, v=50, h=90):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for k, val in (("top", v), ("bottom", v), ("start", h), ("end", h)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def table(headers, rows, widths, head_fill="00594F", size=9.5, code_cols=(), zebra=True, header_size=10):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    tblPr.append(OxmlElement("w:bidiVisual"))
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "D9C89E"); b.append(el)
    tblPr.append(b)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    t.autofit = False
    W = [Cm(w) for w in widths]
    for i, (c, h) in enumerate(zip(t.rows[0].cells, headers)):
        c.width = W[i]; shade(c, head_fill); margins(c)
        p = c.paragraphs[0]; bidi(p); add_text(p, h, header_size, True, WHITE)
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); trPr.append(th)
    for ri, row in enumerate(rows):
        for ci, (c, v) in enumerate(zip(t.rows[ri + 1].cells, row)):
            c.width = W[ci]; margins(c)
            if zebra and ri % 2: shade(c, "F7F4EF")
            p = c.paragraphs[0]; bidi(p); p.paragraph_format.space_after = Pt(0)
            if ci in code_cols: add_code(p, str(v), size - 0.5)
            else: add_text(p, str(v), size)
    para("", after=4)
    return t


def callout(text, fill="EAF4F2", border="8FC7BE", label="ملاحظة", color=GREEN):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr; tblPr.append(OxmlElement("w:bidiVisual"))
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), border); b.append(el)
    tblPr.append(b)
    c = t.rows[0].cells[0]; shade(c, fill); margins(c, 80, 140)
    p = c.paragraphs[0]; bidi(p)
    add_text(p, label + ": ", 10.5, True, color); add_text(p, text, 10.5)
    para("", after=3)


def field(p, instr, placeholder=""):
    for kind, extra in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        r = p.add_run()
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = extra; r._r.append(it)
        elif kind == "text":
            r.text = extra; set_font(r, 10, False, SOFT)
        else:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r._r.append(f)


def constraint_text(t, c, pk):
    bits = []
    if c.name in pk: bits.append("مفتاح أساسي" + (" تلقائي" if c.ai else ""))
    if c.uq: bits.append("فريد")
    tgt = c.fk_target()
    if tgt:
        tt, tc = tgt
        bits.append(f"← {tt}" + ("" if tc == "id" else f".{tc}"))
        if c.od != "RESTRICT": bits.append(f"عند الحذف {c.od}")
    bits.append("اختياري" if c.null else "إلزامي")
    if c.d is not None and not c.onupdate and "CURRENT_TIMESTAMP" not in c.d: bits.append(f"افتراضي {c.d}")
    if c.d is not None and "CURRENT_TIMESTAMP" in c.d: bits.append("الآن افتراضياً")
    return "، ".join(bits)


def render_table(t, n):
    pk = effective_pk(t)
    hp = heading(f"{t.name} — {t.title}", 2, before=14)
    para(t.desc, after=4, line=1.3)
    meta = []
    if t.new: meta.append(("مقترح", "لا مقابل مخزَّن له في النسخة التجريبية"))
    if t.src: meta.append(("المصدر في الكود", t.src))
    if t.seed: meta.append(("البيانات الأولية", t.seed))
    for label, val in meta:
        p = para("", after=2, size=10)
        add_text(p, label + ": ", 10, True, MAROON if label == "مقترح" else GOLD)
        if label == "المصدر في الكود": add_code(p, val.replace("؛", " |").replace("،", ","), 8.5, False, SOFT)
        else: add_text(p, val, 10, False, SOFT)
    rows = [[c.name, c.type.replace(",", ", "), constraint_text(t, c, pk), c.desc] for c in t.cols]
    table(["العمود", "النوع", "القيود", "الوصف"], rows, [4.0, 2.8, 3.0, 7.6], code_cols=(0, 1), size=9)
    keys = []
    if len(pk) > 1: keys.append("المفتاح الأساسي المركب: (" + "، ".join(pk) + ")")
    for u in t.uq:
        if list(u) != pk: keys.append("فريد معاً: (" + "، ".join(u) + ")")
    for ix in t.idx: keys.append("فهرس: (" + "، ".join(ix) + ")")
    for ck in t.checks: keys.append("قيد تحقق: " + ck)
    if keys:
        p = para("المفاتيح والقيود:", size=10.5, bold=True, color=GREEN, after=1)
        bullets(keys, 10)
    if t.rules:
        para("قواعد العمل:", size=10.5, bold=True, color=GREEN, after=1, before=2)
        bullets(t.rules, 10.5)


def render_blocks(blocks):
    for b in blocks:
        kind = b[0]
        if kind == "h": heading(b[1], 2, before=16)
        elif kind == "h3": heading(b[1], 3, before=10)
        elif kind == "p": para(b[1], after=6, line=1.3)
        elif kind == "list": bullets(b[1], 10.5)
        elif kind == "table": table(b[1], b[2], b[3], size=9.5)
        elif kind == "note": callout(b[1])


def build_docx(out_path, n_fks):
    global doc
    doc = Document()
    s0 = doc.sections[0]; s0.page_width, s0.page_height = Cm(21), Cm(29.7)
    for s in doc.sections: s.left_margin = s.right_margin = Cm(1.8); s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(11)
    for lvl, sz, col in ((1, 22, GREEN), (2, 14.5, MAROON), (3, 12.5, GREEN)):
        hs = doc.styles[f"Heading {lvl}"]; hs.font.name = FONT; hs.font.size = Pt(sz); hs.font.color.rgb = col; hs.font.bold = True

    # ── cover ──
    para("", after=36)
    p = doc.add_paragraph(); bidi(p, WD_ALIGN_PARAGRAPH.CENTER); p.add_run().add_picture(str(LOGO), width=Cm(4))
    para("إدارة الحج والعمرة السورية", 14, True, GOLD, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("المنصة الوطنية للحج", 18, True, GREEN, after=30, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(META["title"], 40, True, GREEN, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(META["subtitle"], 20, True, MAROON, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(META["tagline"], 13, False, SOFT, after=36, align=WD_ALIGN_PARAGRAPH.CENTER)
    t = doc.add_table(rows=1, cols=4); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    for c, (a, b) in zip(t.rows[0].cells, [("الجداول", str(len(TABLES))), ("العلاقات", str(n_fks)), ("قاعدة البيانات", "MySQL 8 / MariaDB"), ("تاريخ الإعداد", datetime.date.today().strftime("%d/%m/%Y"))]):
        shade(c, "F7F4EF"); margins(c, 120, 120)
        q = c.paragraphs[0]; bidi(q, WD_ALIGN_PARAGRAPH.CENTER); add_text(q, a, 10, True, GOLD)
        q2 = c.add_paragraph(); bidi(q2, WD_ALIGN_PARAGRAPH.CENTER); add_text(q2, b, 12, True, GREEN)
    para("", after=50)
    para("وثيقة تقنية لفريق التطوير. أسماء الأشخاص والأرقام الواردة أمثلةً مأخوذة من البيانات الوهمية للنسخة التجريبية.", 10, False, SOFT, align=WD_ALIGN_PARAGRAPH.CENTER)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.left_margin = body.right_margin = Cm(1.8); body.top_margin = Cm(1.8); body.bottom_margin = Cm(1.6)
    body.header.is_linked_to_previous = False; body.footer.is_linked_to_previous = False
    hp = body.header.paragraphs[0]; bidi(hp, WD_ALIGN_PARAGRAPH.CENTER)
    add_text(hp, META["title"] + " ", 9.5, True, GREEN); add_text(hp, "— المنصة الوطنية للحج — موسم 1448هـ", 9.5, False, SOFT)
    fp = body.footer.paragraphs[0]; bidi(fp, WD_ALIGN_PARAGRAPH.CENTER)
    add_text(fp, "صفحة ", 9.5, False, SOFT); field(fp, "PAGE", "1"); add_text(fp, "   |   إدارة الحج والعمرة السورية", 9, False, SOFT)
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); body._sectPr.append(pg)

    # ── introduction ──
    heading("مقدمة", 1, before=0)
    for x in INTRO: para(x, after=8, line=1.35)
    heading("كيف تقرأ هذه الوثيقة؟", 2)
    for x in HOW_TO_READ: para(x, after=4)
    bullets(HOW_ITEMS)
    heading("الاصطلاحات المعتمدة في التصميم", 2)
    table(["البند", "الاصطلاح"], [[a, b] for a, b in CONVENTIONS], [3.6, 13.4], size=10)
    heading("خريطة الفصول", 2)
    para("كل فصل مجال، وهذه أهم جداوله وعدد جداوله كلها. وتلتقي الفصول عند جداول قليلة: seasons وaccounts وcivil_persons وapplications وadmin_groups وclusters وstaff_users وfiles.", after=6)
    rows = []
    for i, m in enumerate(MODULES, 1):
        top = "، ".join(t.name for t in m.tables[:6]) + ("…" if len(m.tables) > 6 else "")
        rows.append([str(i), m.title, str(len(m.tables)), top])
    table(["#", "الفصل", "الجداول", "أهمها"], rows, [0.8, 6.0, 1.4, 8.8], size=9.5, code_cols=(3,))
    heading("رحلة البيانات من الحساب إلى العودة", 2)
    bullets([
        "civil_persons ← accounts ← account_season_paths: الشخص موجود في السجل، فيُنشئ حسابه، ويختار مساره في الموسم.",
        "المسار pilgrim: application_drafts ← applications (application_members، payments) ← post_acceptance والخطوات السبع ← member_visas ← tickets وseason_ratings ← hajj_certificates.",
        "المسار admin: admin_season_files ← exam_attempts وoral_bookings ← admin_groups وclusters ← field_records وdaily_reports ← self_evaluations وpeer_evaluations ← grading_results.",
        "يلتقي المساران في post_acceptance.group_id: العائلة تُلحق بمجموعة يقودها إداري، وتركب رحلات مجموعتها (group_flight_legs ← flight_assignments)، وتسكن حيث يسكن تكتلها (op_inner_node_clusters).",
        "والموظفون (staff_users) يقررون ويراجعون ويديرون في كل خطوة، وكل ذلك في audit_events. ومحتوى الصفحات العامة في cms_* وجداول المحتوى.",
    ])

    heading("فهرس المحتويات", 1, page_break=True, before=0)
    tp = doc.add_paragraph(); bidi(tp)
    field(tp, 'TOC \\o "1-2" \\h \\z \\u', "اضغط بزر الفأرة الأيمن هنا ثم «تحديث الحقل» لعرض الفهرس.")

    # ── chapters ──
    for i, m in enumerate(MODULES, 1):
        heading(f"الفصل {i}: {m.title}", 1, page_break=True, before=0)
        for x in m.intro: para(x, after=8, line=1.35)
        if m.flow:
            para("كيف تتدفق البيانات:", 12, True, MAROON, after=2, before=4)
            bullets(m.flow, 11, style="List Number")
        para(f"جداول الفصل ({len(m.tables)}): " + "، ".join(t.name for t in m.tables), 10, False, SOFT, after=6)
        for n, t in enumerate(m.tables, 1): render_table(t, n)
        rels = []
        for t in m.tables:
            for c in t.cols:
                tgt = c.fk_target()
                if not tgt: continue
                tt, tc = tgt
                one = c.uq or [c.name] == effective_pk(t)
                rels.append([f"{t.name}.{c.name}", f"{tt}.{tc}", "واحد إلى واحد" if one else "متعدد إلى واحد", c.od])
        if rels:
            heading(f"علاقات الفصل ({len(rels)})", 2, before=16)
            table(["من", "إلى", "العلاقة", "عند الحذف"], rels, [7.0, 5.2, 2.7, 2.5], size=9, code_cols=(0, 1, 3))
        render_blocks(m.blocks)

    # ── appendices ──
    heading("ملحق 1: فجوات النسخة التجريبية وما يقابلها في هذا التصميم", 1, page_break=True, before=0)
    para("ما وجدناه في كود النسخة التجريبية ولا يصلح لخادم حقيقي، وكيف عالجه التصميم. كل بند منها يحتاج أيضاً تعديلاً في واجهة المنصة حين تُربط بالخادم.", after=6)
    bullets(DEMO_GAPS, 10.5)
    heading("ملحق 2: فهرس الجداول", 1, page_break=True, before=0)
    idx = sorted(((t.name, t.title, next(i for i, m in enumerate(MODULES, 1) if t in m.tables)) for t in TABLES))
    table(["الجدول", "الاسم", "الفصل"], [[a, b, str(c)] for a, b, c in idx], [6.0, 9.4, 1.6], size=9, code_cols=(0,))
    heading("ملحق 3: كيف تُولَّد هذه الوثيقة وملف SQL", 1, page_break=True, before=0)
    bullets([
        "المصدر: docs/_build-guide/schema/ — ملف لكل فصل (m01_core.py … m16_cms.py) بلغة وصف بسيطة (schema/dsl.py): T للجدول، c للعمود، وعلاقاته وقيوده وقواعده.",
        "التوليد: python build_schema.py <المجلد> يتحقق من النموذج (كل مفتاح أجنبي يشير إلى مفتاح موجود بنوع مطابق، ولكل جدول مفتاح أساسي، وأعمدة كل مفتاح موجودة) ثم يكتب الوثيقة وschema.sql.",
        "ثم: powershell -File finalize.ps1 -Docx <الوثيقة> -Pdf <PDF> لتحديث الفهرس والطباعة (يحتاج Word على Windows).",
        "التطبيق: mysql -u root -p hajj < schema.sql على قاعدة فارغة بترميز utf8mb4. الجداول تُنشأ أولاً ثم تُضاف العلاقات في آخر الملف.",
        "أي تغيير في المنصة يمس بياناتها يُعدَّل في النموذج، ويُعاد توليد الوثيقة والملف معاً.",
    ], 10.5)

    # ── RTL everywhere ──
    dd = doc.styles.element.find(qn("w:docDefaults"))
    ppd = dd.find(qn("w:pPrDefault"))
    if ppd is None: ppd = OxmlElement("w:pPrDefault"); dd.append(ppd)
    ppr = ppd.find(qn("w:pPr"))
    if ppr is None: ppr = OxmlElement("w:pPr"); ppd.append(ppr)
    if ppr.find(qn("w:bidi")) is None: ppr.insert(0, OxmlElement("w:bidi"))
    from docx.enum.style import WD_STYLE_TYPE
    for nm in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "toc 1", "toc 2"]:
        try: s = doc.styles[nm]
        except KeyError: s = doc.styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
        pPr = s.element.get_or_add_pPr()
        if pPr.find(qn("w:bidi")) is None: pPr.insert(0, OxmlElement("w:bidi"))
        r = s.element.get_or_add_rPr()
        rf = r.find(qn("w:rFonts"))
        if rf is None: rf = OxmlElement("w:rFonts"); r.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
        if nm.startswith("toc"):
            s.base_style = doc.styles["Normal"]; s.font.size = Pt(11 if nm == "toc 1" else 10); s.font.bold = nm == "toc 1"
    for sct in doc.sections:
        sp = sct._sectPr
        if sp.find(qn("w:bidi")) is None: sp.append(OxmlElement("w:bidi"))
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); doc.settings.element.append(uf)
    _TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
    for tb in doc.tables:
        tp_ = tb._tbl.tblPr; kids = list(tp_)
        for e in kids: tp_.remove(e)
        for e in sorted(kids, key=lambda e: _TBLPR.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in _TBLPR else 99): tp_.append(e)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


if __name__ == "__main__":
    errors, warns = validate()
    for w_ in warns: print("warn:", w_)
    for e in errors: print("ERROR:", e)
    print(f"{len(MODULES)} chapters, {len(TABLES)} tables, {sum(len(t.cols) for t in TABLES)} columns, {len(errors)} errors, {len(warns)} warnings")
    if errors: sys.exit(1)
    if sys.argv[1:] == ["--check"]: sys.exit(0)
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT.parent / "هيكلية-قاعدة-البيانات"
    sql, n_fks = build_sql()
    (out).mkdir(parents=True, exist_ok=True)
    (out / "schema.sql").write_text(sql, "utf-8")
    build_docx(out / "هيكلية-قاعدة-البيانات.docx", n_fks)
    print("saved", out / "schema.sql", "and", out / "هيكلية-قاعدة-البيانات.docx", "| foreign keys:", n_fks)
