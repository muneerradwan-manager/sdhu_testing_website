# -*- coding: utf-8 -*-
import json, sys, datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import importlib
import re
from docx.text.paragraph import Paragraph
_LTR = re.compile(r"\+?[A-Za-z]*\d[\dA-Za-z]*(?:[ \-][A-Za-z]*\d[\dA-Za-z]*)+|[A-Z]{2,}-\d+|\d+-[A-Z]-\d+(?:-\d+)?")
import copy

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT))
C = importlib.import_module(sys.argv[1])
SHOTS = Path(sys.argv[2]) if Path(sys.argv[2]).is_absolute() else ROOT / sys.argv[2]
OUT = Path(sys.argv[3])
M = C.META
FONT = "itf Qomra Arabic"
GREEN = RGBColor(0x00, 0x59, 0x4F); GREEN2 = RGBColor(0x28, 0x9E, 0x92)
GOLD = RGBColor(0xAD, 0x9E, 0x6E); MAROON = RGBColor(0x67, 0x21, 0x46)
INK = RGBColor(0x02, 0x15, 0x26); SOFT = RGBColor(0x55, 0x5F, 0x6B)
# The site's own icon, found from this folder (docs/_build-guide) on any machine
LOGO = ROOT.parent.parent / "public" / "icons" / "icon-512.png"

man = json.loads((SHOTS / "manifest.json").read_text("utf-8"))
IMG = {}
for e in man:
    IMG[e["key"]] = (SHOTS / "jpg" / (Path(e["file"]).stem + ".jpg"), e["caption"])

doc = Document()

# ── page setup ──
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2); s.top_margin = Cm(2); s.bottom_margin = Cm(1.8)

# ── helpers ──
def set_run_font(run, size=None, bold=None, color=None, italic=None, font=FONT):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.insert(0, rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), font)
    if size is not None:
        run.font.size = Pt(size)
        szCs = rPr.find(qn("w:szCs"))
        if szCs is None:
            szCs = OxmlElement("w:szCs"); rPr.append(szCs)
        szCs.set(qn("w:val"), str(int(size * 2)))
    if bold is not None:
        run.font.bold = bold
        bCs = rPr.find(qn("w:bCs"))
        if bCs is None:
            bCs = OxmlElement("w:bCs"); rPr.append(bCs)
        bCs.set(qn("w:val"), "1" if bold else "0")
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rtl = rPr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl"); rPr.append(rtl)
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rPr.append(lang)
    lang.set(qn("w:bidi"), "ar-SY")
    _split_ltr(run)

_PAREN_OPEN = re.compile(r"\((?=[\d%])")
_PAREN_CLOSE = re.compile(r"(?<=[\d%])\)")
def _split_ltr(run):
    text = _PAREN_CLOSE.sub("‏)", _PAREN_OPEN.sub("(‏", run.text))
    run.text = text
    parts, pos = [], 0
    for m in _LTR.finditer(text):
        if m.start() > pos: parts.append((text[pos:m.start()], False))
        parts.append((m.group(0), True)); pos = m.end()
    if pos < len(text): parts.append((text[pos:], False))
    if not parts or (len(parts) == 1 and not parts[0][1]): return
    template = copy.deepcopy(run._r)
    run.text = parts[0][0]
    if parts[0][1]:
        for e in run._r.rPr.findall(qn("w:rtl")): run._r.rPr.remove(e)
    prev = run._r
    for t, ltr in parts[1:]:
        r2 = copy.deepcopy(template)
        for ch in list(r2):
            if ch.tag != qn("w:rPr"): r2.remove(ch)
        te = OxmlElement("w:t"); te.set(qn("xml:space"), "preserve"); te.text = t; r2.append(te)
        if ltr:
            for e in r2.rPr.findall(qn("w:rtl")): r2.rPr.remove(e)
        prev.addnext(r2); prev = r2

def set_rtl(p, align=WD_ALIGN_PARAGRAPH.LEFT):
    # In a bidi paragraph Word reads jc="left" as START (the right side) — never use RIGHT here
    pPr = p._p.get_or_add_pPr()
    p.alignment = align
    if pPr.find(qn("w:bidi")) is None:
        b = OxmlElement("w:bidi"); j = pPr.find(qn("w:jc"))
        pPr.insert(list(pPr).index(j), b) if j is not None else pPr.append(b)
    return p

def para(container, text="", size=12, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, after=6, before=0, style=None, italic=False, line=1.15):
    p = container.add_paragraph(style=style) if style else container.add_paragraph()
    set_rtl(p, align)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = line
    if text:
        r = p.add_run(text); set_run_font(r, size, bold, color, italic)
    return p

def add_runs(p, parts):
    """parts: list of (text, dict(size,bold,color))"""
    for text, kw in parts:
        r = p.add_run(text); set_run_font(r, kw.get("size", 12), kw.get("bold", False), kw.get("color", INK), kw.get("italic"))

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def cell_margins(cell, top=80, bottom=80, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", right), ("end", left)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

_TBLPR_ORDER = ["tblStyle","tblpPr","tblOverlap","bidiVisual","tblStyleRowBandSize","tblStyleColBandSize","tblW","jc","tblCellSpacing","tblInd","tblBorders","shd","tblLayout","tblCellMar","tblLook","tblCaption","tblDescription"]
def tidy(t):
    tblPr = t._tbl.tblPr
    kids = list(tblPr)
    def key(e):
        tag = e.tag.split("}")[1]
        return _TBLPR_ORDER.index(tag) if tag in _TBLPR_ORDER else 99
    for e in kids: tblPr.remove(e)
    for e in sorted(kids, key=key): tblPr.append(e)

def fix_widths(t, widths):
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tblPr = t._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(sum(w for w in widths) / 635)))
    for col, w in zip(t.columns, widths):
        col.width = w
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = w

def table_rtl(t):
    tblPr = t._tbl.tblPr
    bv = OxmlElement("w:bidiVisual"); tblPr.append(bv)

def set_borders(t, color="D9C89E", sz=6, inside=True):
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for e in edges:
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), color); b.append(el)
    tblPr.append(b)

def no_borders(t):
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "nil"); b.append(el)
    tblPr.append(b)

def callout(kind, text):
    cfg = {"warn": ("تنبيه", "FBEFF3", MAROON), "note": ("ملاحظة", "EAF4F2", GREEN), "tip": ("نصيحة", "F7F1DF", GOLD)}[kind]
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; table_rtl(t)
    set_borders(t, color={"warn": "C98AA5", "note": "8FC7BE", "tip": "D9C89E"}[kind], sz=8, inside=False)
    c = t.rows[0].cells[0]; shade(c, cfg[1]); cell_margins(c)
    p = c.paragraphs[0]; set_rtl(p); p.paragraph_format.space_after = Pt(2)
    add_runs(p, [(cfg[0] + ": ", {"bold": True, "color": cfg[2], "size": 11.5}), (text, {"size": 11.5})])
    para(doc, "", after=4)

def heading(text, level, color=GREEN, size=None, page_break=False, before=12):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_rtl(p)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(6)
    if page_break: p.paragraph_format.page_break_before = True
    r = p.add_run(text); set_run_font(r, size or {1: 22, 2: 17, 3: 13.5}[level], True, color)
    return p

def bullets(items, size=12):
    for it in items:
        p = doc.add_paragraph(style="List Bullet"); set_rtl(p)
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.15
        r = p.add_run(it); set_run_font(r, size, False, INK)

def picture(key, fig_no):
    path, cap = IMG[key]
    p = doc.add_paragraph(); set_rtl(p); p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(16.6))
    cp = para(doc, f"صورة {fig_no}: {cap}", size=10, color=SOFT, after=8, italic=False)
    return cp

def field(p, instr, placeholder=""):
    r = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r2 = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2._r.append(it)
    r3 = p.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate"); r3._r.append(f2)
    r4 = p.add_run(placeholder); set_run_font(r4, 10, False, SOFT)
    r5 = p.add_run(); f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end"); r5._r.append(f3)

# ── base styles ──
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(12)
rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
if rf is None: rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
for lvl, sz, col in ((1, 22, GREEN), (2, 17, MAROON), (3, 13.5, GREEN)):
    hs = doc.styles[f"Heading {lvl}"]; hs.font.name = FONT; hs.font.size = Pt(sz); hs.font.color.rgb = col; hs.font.bold = True
    hr = hs.element.get_or_add_rPr(); hf = hr.find(qn("w:rFonts"))
    if hf is None: hf = OxmlElement("w:rFonts"); hr.insert(0, hf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): hf.set(qn(a), FONT)

# ═══════════════ COVER ═══════════════
today = datetime.date.today()
para(doc, "", after=40)
p = doc.add_paragraph(); set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER); p.add_run().add_picture(str(LOGO), width=Cm(4.2))
para(doc, "إدارة الحج والعمرة السورية", size=14, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, bold=True)
para(doc, "المنصة الوطنية للحج", size=18, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER, after=30, bold=True)
para(doc, M["title"], size=44, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, bold=True)
para(doc, M["subtitle"], size=22, color=MAROON, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, bold=True)
para(doc, M["tagline"], size=14, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; table_rtl(t); no_borders(t)
for c, (a, b) in zip(t.rows[0].cells, [("موسم الحج", "1448هـ — 2027م"), ("الإصدار", "النسخة التجريبية (Demo)"), ("تاريخ الإعداد", today.strftime("%d/%m/%Y"))]):
    shade(c, "F7F4EF"); cell_margins(c, 120, 120)
    q = c.paragraphs[0]; set_rtl(q, WD_ALIGN_PARAGRAPH.CENTER); r = q.add_run(a); set_run_font(r, 10.5, True, GOLD)
    q2 = c.add_paragraph(); set_rtl(q2, WD_ALIGN_PARAGRAPH.CENTER); r = q2.add_run(b); set_run_font(r, 12, True, GREEN)
para(doc, "", after=60)
para(doc, "جميع الأسماء والأرقام والبيانات الواردة في هذا الدليل وهمية، مأخوذة من النسخة التجريبية للمنصة لأغراض العرض والتدريب.", size=10.5, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

# ═══════════════ BODY SECTION (own header/footer) ═══════════════
body = doc.add_section(WD_SECTION.NEW_PAGE)
body.left_margin = body.right_margin = Cm(2); body.top_margin = Cm(2); body.bottom_margin = Cm(1.8)
body.header.is_linked_to_previous = False; body.footer.is_linked_to_previous = False
# cover section: empty header/footer
for s in (doc.sections[0],):
    s.header.paragraphs[0].text = ""; s.footer.paragraphs[0].text = ""
hp = body.header.paragraphs[0]; set_rtl(hp, WD_ALIGN_PARAGRAPH.CENTER)
add_runs(hp, [(M["title"] + " ", {"bold": True, "color": GREEN, "size": 10}), ("— المنصة الوطنية للحج — موسم 1448هـ — نسخة تجريبية", {"color": SOFT, "size": 10})])
pb = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom"); bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "6"); bt.set(qn("w:space"), "1"); bt.set(qn("w:color"), "D9C89E"); pb.append(bt); hp._p.get_or_add_pPr().append(pb)
fp = body.footer.paragraphs[0]; set_rtl(fp, WD_ALIGN_PARAGRAPH.CENTER)
r = fp.add_run("صفحة "); set_run_font(r, 10, False, SOFT); field(fp, "PAGE", "1")
r = fp.add_run("   |   جميع البيانات وهمية لأغراض العرض   |   إدارة الحج والعمرة السورية"); set_run_font(r, 9.5, False, SOFT)
# page numbering restarts at 1
sectPr = body._sectPr; pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); sectPr.append(pg)

# ── مقدمة ──
heading("مقدمة", 1, before=0)
for t_ in C.INTRO: para(doc, t_, after=8, line=1.3)
if M.get("site_url"):
    from site_link import add_to_guide
    add_to_guide(doc, M, C.INTRO)
heading("كيف تقرأ هذا الدليل؟", 2)
for t_ in C.HOW_TO_READ: para(doc, t_, after=4, line=1.3)
bullets(C.HOW_ITEMS)
heading("ما يجب أن تعرفه عن النسخة التجريبية", 2)
bullets(C.DEMO_NOTES)

heading(M["journey_title"], 2)
para(doc, M["journey_intro"], after=6)
t = doc.add_table(rows=len(C.JOURNEY), cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; table_rtl(t); set_borders(t, "FFFFFF", 12)
widths = [Cm(1.4), Cm(4.2), Cm(11.4)]
fix_widths(t, widths)
for i, (n, name, desc) in enumerate(C.JOURNEY):
    cells = t.rows[i].cells
    for c, w in zip(cells, widths): c.width = w
    shade(cells[0], "00594F"); shade(cells[1], "EAF4F2"); shade(cells[2], "F7F4EF")
    for c in cells: cell_margins(c, 90, 90)
    q = cells[0].paragraphs[0]; set_rtl(q, WD_ALIGN_PARAGRAPH.CENTER); r = q.add_run(n); set_run_font(r, 16, True, RGBColor(0xD9, 0xC8, 0x9E))
    q = cells[1].paragraphs[0]; set_rtl(q); r = q.add_run(name); set_run_font(r, 12.5, True, GREEN)
    q = cells[2].paragraphs[0]; set_rtl(q); r = q.add_run(desc); set_run_font(r, 11, False, INK)
para(doc, "", after=4)

# ── فهرس ──
heading("فهرس المحتويات", 1, page_break=True, before=0)
tp = doc.add_paragraph(); set_rtl(tp)
field(tp, 'TOC \\o "1-2" \\h \\z \\u', "اضغط بزر الفأرة الأيمن هنا ثم «تحديث الحقل» لعرض الفهرس.")

# ── الفصول ──
fig = 0
for ci, ch in enumerate(C.CHAPTERS, 1):
    heading(f"الفصل {ci}: {ch['title']}", 1, page_break=True, before=0)
    para(doc, ch["intro"], after=10, line=1.3)
    step_no = 0
    for sec_ in ch["sections"]:
        heading(sec_["title"], 2)
        for s in sec_["steps"]:
            step_no += 1
            hp_ = heading(f"الخطوة {step_no}: {s['t']}", 3, before=10)
            hp_.paragraph_format.keep_with_next = True
            para(doc, s["x"], after=6, line=1.3).paragraph_format.keep_with_next = True
            if s["img"]:
                fig += 1; picture(s["img"], fig)
            s["do"] = [d for d in s["do"] if d.strip() not in ("—", "")]
            if s["do"]:
                para(doc, "ما يجب عليك فعله:", size=12, bold=True, color=MAROON, after=2).paragraph_format.keep_with_next = True
                bullets(s["do"])
            if s["r"] and s["r"].strip() != "—":
                p = para(doc, "", after=8, before=4); add_runs(p, [("النتيجة المتوقعة: ", {"bold": True, "color": GREEN}), (s["r"], {})])
            if s["warn"]: callout("warn", s["warn"])
            if s["note"]: callout("note", s["note"])
            if s["tip"]: callout("tip", s["tip"])

# ── الأسئلة الشائعة ──
heading("الأسئلة والمشكلات الشائعة", 1, page_break=True, before=0)
para(doc, M["faq_intro"], after=10, line=1.3)
for q, a in C.FAQ:
    p = para(doc, q, size=12.5, bold=True, color=GREEN, after=2, before=6); p.paragraph_format.keep_with_next = True
    para(doc, a, after=8, line=1.3)

# ── ملاحق ──
def simple_table(headers, rows, widths, head_fill="00594F"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; table_rtl(t); set_borders(t); fix_widths(t, widths)
    for c, h, w in zip(t.rows[0].cells, headers, widths):
        c.width = w; shade(c, head_fill); cell_margins(c)
        q = c.paragraphs[0]; set_rtl(q); r = q.add_run(h); set_run_font(r, 11.5, True, RGBColor(0xFF, 0xFF, 0xFF))
    for i, row in enumerate(rows):
        for c, v, w in zip(t.rows[i + 1].cells, row, widths):
            c.width = w; cell_margins(c)
            if i % 2: shade(c, "F7F4EF")
            q = c.paragraphs[0]; set_rtl(q); r = q.add_run(v); set_run_font(r, 11, False, INK)
    para(doc, "", after=6)

for ai, ap in enumerate(C.APPENDICES, 1):
    heading(f"ملحق {ai}: {ap['title']}", 1, page_break=(ai == 1), before=0 if ai == 1 else 12)
    if ap.get("intro"): para(doc, ap["intro"], after=6, line=1.3)
    simple_table(ap["headers"], ap["rows"], [Cm(w) for w in ap["widths"]])
    if ap.get("note"): para(doc, ap["note"], after=8, line=1.3)

# updateFields on open (fallback if COM update fails)
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)

OUT.parent.mkdir(parents=True, exist_ok=True)

# ═══════════════ RTL at every level: defaults, styles, TOC styles, sections ═══════════════
from docx.enum.style import WD_STYLE_TYPE
_PPR_AFTER_BIDI = ["adjustRightInd","snapToGrid","spacing","ind","contextualSpacing","mirrorIndents","suppressOverlap","jc","textDirection","textAlignment","textboxTightWrap","outlineLvl","divId","cnfStyle","rPr","sectPr","pPrChange"]
def ppr_bidi(pPr):
    if pPr.find(qn("w:bidi")) is not None: return
    b = OxmlElement("w:bidi")
    for i, ch in enumerate(list(pPr)):
        if ch.tag.split("}")[1] in _PPR_AFTER_BIDI:
            pPr.insert(i, b); return
    pPr.append(b)
def style_rtl(name, jc="right"):
    try: s = doc.styles[name]
    except KeyError: s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    pPr = s.element.get_or_add_pPr(); ppr_bidi(pPr)
    if jc:
        j = pPr.find(qn("w:jc"))
        if j is None: j = OxmlElement("w:jc"); pPr.append(j)
        j.set(qn("w:val"), jc)
    rPr = s.element.get_or_add_rPr()
    if rPr.find(qn("w:rtl")) is None: rPr.append(OxmlElement("w:rtl"))
    rf = rPr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    return s
# document defaults
dd = doc.styles.element.find(qn("w:docDefaults"))
ppd = dd.find(qn("w:pPrDefault"))
if ppd is None: ppd = OxmlElement("w:pPrDefault"); dd.append(ppd)
ppr = ppd.find(qn("w:pPr"))
if ppr is None: ppr = OxmlElement("w:pPr"); ppd.append(ppr)
ppr_bidi(ppr)
for nm in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "Header", "Footer", "Caption", "Title"]:
    style_rtl(nm, None if nm in ("Header", "Footer") else "left")
for nm in ["toc 1", "toc 2", "toc 3"]:
    s = style_rtl(nm, "left"); s.base_style = doc.styles["Normal"]
    if nm == "toc 1": s.font.bold = True
    s.font.size = Pt(12 if nm == "toc 1" else 11)
    s.paragraph_format.space_after = Pt(4)
# every paragraph (incl. table cells, header/footer) gets bidi in the right position
for pel in doc.element.body.iter(qn("w:p")):
    ppr_bidi(pel.get_or_add_pPr())
for sct in doc.sections:
    for part in (sct.header, sct.footer):
        for pel in part._element.iter(qn("w:p")): ppr_bidi(pel.get_or_add_pPr())
    sp = sct._sectPr
    if sp.find(qn("w:bidi")) is None:
        b = OxmlElement("w:bidi"); dg = sp.find(qn("w:docGrid"))
        sp.insert(list(sp).index(dg), b) if dg is not None else sp.append(b)
    if sp.find(qn("w:rtlGutter")) is None:
        g = OxmlElement("w:rtlGutter"); dg = sp.find(qn("w:docGrid"))
        sp.insert(list(sp).index(dg), g) if dg is not None else sp.append(g)

for t_ in doc.tables: tidy(t_)
doc.save(str(OUT))
print("saved", OUT, "figures:", fig, "size MB:", round(OUT.stat().st_size / 1e6, 1))
